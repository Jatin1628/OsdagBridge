# if openpyxl not yet installed, run "pip install openpyxl"
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PySide6.QtWidgets import (
    QDialog, QWidget, QLabel, QPushButton, QVBoxLayout, QHBoxLayout,
    QTreeWidget, QTreeWidgetItem, QSizeGrip, QFileDialog,
    QFrame, QTableWidget, QTableWidgetItem, QHeaderView,
    QStyledItemDelegate, QSizePolicy
)
from PySide6.QtCore import Qt, QRect
from PySide6.QtGui import QColor, QFont
from osdagbridge.desktop.ui.utils.custom_titlebar import CustomTitleBar


class CheckboxDelegate(QStyledItemDelegate):
    """Custom delegate to paint checkmarks in checkboxes."""

    def paint(self, painter, option, index):
        super().paint(painter, option, index)

        tree = self.parent()
        if not isinstance(tree, QTreeWidget):
            return

        item = tree.itemFromIndex(index)
        if not item:
            return

        check_state = item.checkState(0)

        if check_state == Qt.Checked:
            symbol = "✓"
        elif check_state == Qt.PartiallyChecked:
            symbol = "−"
        else:
            return

        checkbox_rect = QRect(
            option.rect.left() + 6,
            option.rect.top(),
            20,
            option.rect.height()
        )

        painter.setPen(QColor("#90AF13"))
        painter.setFont(QFont("Arial", 10))
        painter.drawText(checkbox_rect, Qt.AlignCenter, symbol)


class ExportTableDialog(QDialog):
    """Export Results Table Dialog"""

    def __init__(self, selected_tables=None, parent=None):
        super().__init__(parent)

        self.selected_tables = selected_tables or {}
        self.setMinimumWidth(920)
        self.setMinimumHeight(650)
        self.setObjectName("export_table_dialog")

        self.setStyleSheet("""
            QDialog#export_table_dialog {
                background: white;
                border: 1px solid #90AF13;
            }
        """)

        self._is_expanded = True
        self._setup_ui()

    #WRAPPER
    def setupWrapper(self):
        self.setWindowFlags(Qt.FramelessWindowHint | Qt.WindowSystemMenuHint)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(1, 1, 1, 1)
        main_layout.setSpacing(0)

        self.title_bar = CustomTitleBar()
        self.title_bar.setTitle("Export Results")
        main_layout.addWidget(self.title_bar)

        self.content_widget = QWidget(self)
        main_layout.addWidget(self.content_widget, 1)

        size_grip = QSizeGrip(self)
        size_grip.setFixedSize(16, 16)

        overlay = QHBoxLayout()
        overlay.setContentsMargins(0, 0, 4, 4)
        overlay.addStretch()
        overlay.addWidget(size_grip, 0, Qt.AlignBottom | Qt.AlignRight)

        main_layout.addLayout(overlay)

    #MAIN UI
    def _setup_ui(self):
        self.setupWrapper()

        main_layout = QVBoxLayout(self.content_widget)
        main_layout.setContentsMargins(18, 18, 18, 14)
        main_layout.setSpacing(14)

        card = QFrame()
        card_layout = QVBoxLayout(card)
        card_layout.setSpacing(10)

        #MAIN SPLIT
        main_split = QHBoxLayout()
        main_split.setSpacing(0)
        main_split.setContentsMargins(0, 0, 0, 0)

        #LEFT PANEL
        self.left_box = QFrame()
        self.left_box.setObjectName("leftBox")
        self.left_box.setMinimumWidth(320)
        self.left_box.setMaximumWidth(400)
        self.left_box.setStyleSheet("""
            QFrame#leftBox {
                border: 1px solid #90AF13;
                border-radius: 12px;
                background: white;
            }
        """)

        left_layout = QVBoxLayout(self.left_box)
        left_layout.setContentsMargins(8, 10, 8, 8)
        left_layout.setSpacing(8)

        #SELECT / CLEAR ALL buttons
        action_btn_style = """
            QPushButton {
                background: white;
                color: #3a3a3a;
                border: 1px solid #90af13;
                border-radius: 6px;
                padding: 4px 10px;
                font-size: 11px;
                font-weight: 500;
            }
            QPushButton:hover {
                background: #90AF13;
                border-color: #90AF13;
                color: white;
            }
            QPushButton:pressed {
                background: #7a9a10;
                border-color: #7a9a10;
                color: white;
            }
        """

        top_bar = QHBoxLayout()
        top_bar.setSpacing(6)

        title_lbl = QLabel("Export Tables")
        title_lbl.setStyleSheet("""
            QLabel {
                font-size: 13px;
                font-weight: 700;
                color: #2d2d2d;
                border: none;
                background: transparent;
                padding-left: 4px;
            }
        """)
        top_bar.addWidget(title_lbl)
        top_bar.addStretch()

        self.select_btn = QPushButton("Select All")
        self.select_btn.setStyleSheet(action_btn_style)
        self.select_btn.clicked.connect(self.select_all_items)

        self.clear_btn = QPushButton("Clear All")
        self.clear_btn.setStyleSheet(action_btn_style)
        self.clear_btn.clicked.connect(self.clear_all_items)

        top_bar.addWidget(self.select_btn)
        top_bar.addWidget(self.clear_btn)

        left_layout.addLayout(top_bar)

        #TREE
        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.setStyleSheet("""
            QTreeWidget {
                border: none;
                background: white;
                outline: none;
                font-size: 12px;
                color: black;
            }
            QTreeWidget::item {
                height: 28px;
                padding-left: 6px;
                border: none;
            }
            QTreeWidget::item:hover {
                background: rgba(144, 175, 19, 45);
            }
            QTreeWidget::item:selected {
                background: rgba(144, 175, 19, 60);
                color: black;
            }
            QTreeWidget::indicator {
                width: 14px;
                height: 14px;
            }
            QTreeWidget::indicator:unchecked {
                background-color: white;
                border: 1px solid #90AF13;
                border-radius: 3px;
            }
            QTreeWidget::indicator:checked {
                background-color: white;
                border: 1px solid #90AF13;
                border-radius: 3px;
            }
            QTreeWidget::indicator:indeterminate {
                background-color: white;
                border: 1px solid #90AF13;
                border-radius: 3px;
            }
            QScrollBar:vertical {
                width: 2px;
                background: transparent;
                border-radius: 6px;
            }

            QScrollBar:horizontal {
                height: 2px;
                background: transparent;
                border-radius: 6px;
            }

            QScrollBar::handle {
                background: transparent;
                border-radius: 3px;
            }

            QScrollBar::handle:hover,
            QScrollBar::handle:pressed {
                background: #90AF13;
            }

            QScrollBar::add-line,
            QScrollBar::sub-line {
                width: 0px;
                height: 0px;
            }

            QScrollBar::add-page,
            QScrollBar::sub-page {
                background: transparent;
            }
        """)
        self.tree.setItemDelegate(CheckboxDelegate(self.tree))
        self.tree.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.tree.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.tree.setHorizontalScrollMode(QTreeWidget.ScrollPerPixel)
        self.tree.setWordWrap(False)
        self.tree.setUniformRowHeights(False)
        self.tree.setIndentation(20)
        self.tree.itemClicked.connect(self.load_selected_table)
        self.tree.itemChanged.connect(self.handle_item_changed)

        left_layout.addWidget(self.tree, 1)

        #LEFT PANEL CONTAINER
        self.left_container = QWidget()
        self.left_container.setSizePolicy(
            QSizePolicy.Preferred,
            QSizePolicy.Expanding
        )
        left_container_layout = QHBoxLayout(self.left_container)
        left_container_layout.setContentsMargins(0, 0, 8, 0)
        left_container_layout.setSpacing(0)
        left_container_layout.addWidget(self.left_box, 1)

        self._toggle_btn_style_expanded = """
            QPushButton {
                background-color: #90AF13;
                color: white;
                border: none;
                border-top-right-radius: 8px;
                border-bottom-right-radius: 8px;
                border-top-left-radius: 0px;
                border-bottom-left-radius: 0px;
                font-size: 18px;
                font-weight: bold;
                padding: 0px;
                margin: 0px;
            }
            QPushButton:hover { background-color: #7a9a10; }
            QPushButton:pressed { background-color: #6a8a0e; }
        """
        self._toggle_btn_style_collapsed = """
            QPushButton {
                background-color: #90AF13;
                color: white;
                border: none;
                border-radius: 8px;
                font-size: 18px;
                font-weight: bold;
                padding: 0px;
                margin: 0px;
            }
            QPushButton:hover { background-color: #7a9a10; }
            QPushButton:pressed { background-color: #6a8a0e; }
        """

        self.toggle_btn = QPushButton("‹")
        self.toggle_btn.setFixedSize(20, 56)
        self.toggle_btn.setCursor(Qt.PointingHandCursor)
        self.toggle_btn.clicked.connect(self.toggle_tree_panel)
        self.toggle_btn.setStyleSheet(self._toggle_btn_style_expanded)
        left_container_layout.addWidget(self.toggle_btn, 0, Qt.AlignVCenter)

        main_split.addWidget(self.left_container)

        #RIGHT PANEL
        self.right_box = QFrame()
        self.right_box.setAttribute(Qt.WA_StyledBackground, True)
        self.right_box.setStyleSheet("""
            QFrame {
                border: 1px solid #90AF13;
                border-radius: 12px;
                background: white;
            }
        """)

        right_layout = QVBoxLayout(self.right_box)
        right_layout.setContentsMargins(0, 0, 0, 3)
        right_layout.setSpacing(0)

        self.preview = QTableWidget()
        self.preview.viewport().setStyleSheet("""
            background: white;
            border-bottom-left-radius: 11px;
            border-bottom-right-radius: 11px;
        """)
        self.preview.setColumnCount(2)
        self.preview.setHorizontalHeaderLabels(["Parameter", "Value"])
        self.preview.horizontalHeader().setSectionResizeMode(QHeaderView.Fixed)
        self.preview.horizontalHeader().setSectionsClickable(False)
        self.preview.horizontalHeader().setStretchLastSection(True)
        self.preview.verticalHeader().setVisible(False)
        self.preview.setFrameShape(QFrame.NoFrame)
        self.preview.setShowGrid(True)
        self.preview.setStyleSheet("""
            QTableWidget {
                border: none;
                background: white;
                border-bottom-left-radius: 11px;
                border-bottom-right-radius: 11px;
            }
            QTableWidget::item {
                padding: 4px;
                color: black;
                background: white;
            }
            QTableWidget::item:selected {
                color: black;
                background: rgba(144, 175, 19, 60);
            }
            QHeaderView::section {
                background: #90AF13;
                color: white;
                border: none;
                padding: 6px;
                font-weight: 600;
            }
            QHeaderView {
                border: none;
            }
            QHeaderView::section:first {
                border-top-left-radius: 11px;
            }
            QHeaderView::section:last {
                border-top-right-radius: 11px;
            }
            QTableCornerButton::section {
                border: none;
                background: #90AF13;
            }
            QScrollBar:vertical {
                width: 4px;
                background: transparent;
            }
            QScrollBar::handle:vertical {
                background: #90AF13;
                min-height: 20px;
                border-radius: 2px;
            }
            QScrollBar::add-line:vertical,
            QScrollBar::sub-line:vertical {
                height: 0px;
            }
            QScrollBar::add-page:vertical,
            QScrollBar::sub-page:vertical {
                background: transparent;
            }
        """)

        right_layout.addWidget(self.preview)
        main_split.addWidget(self.right_box, 1)

        card_layout.addLayout(main_split)
        main_layout.addWidget(card, 1)

        #FOOTER
        footer = QHBoxLayout()
        footer.addStretch()

        btn_style = """
            QPushButton {
                background: white;
                color: black;
                border: 1px solid #c0c0c0;
                border-radius: 6px;
                padding: 8px 18px;
                font-weight: 600;
                min-width: 110px;
            }
            QPushButton:hover {
                background: #90AF13;
                border-color: #90AF13;
                color: white;
            }
        """

        cancel_btn = QPushButton("Cancel")
        cancel_btn.setStyleSheet(btn_style)
        cancel_btn.clicked.connect(self.reject)

        export_btn = QPushButton("Export")
        export_btn.setStyleSheet(btn_style)
        export_btn.clicked.connect(self.export_excel)

        footer.addWidget(cancel_btn)
        footer.addWidget(export_btn)

        main_layout.addLayout(footer)

        self.populate_tree()


    #LEFT PANEL TOGGLE
    def toggle_tree_panel(self):
        if self._is_expanded:
            self._is_expanded = False
            self.left_box.setFixedWidth(70)
            self.toggle_btn.setText("›")
            self.toggle_btn.setStyleSheet(self._toggle_btn_style_collapsed)
            self.select_btn.hide()
            self.clear_btn.hide()
        else:
            self._is_expanded = True
            self.left_box.setMinimumWidth(272)
            self.left_box.setMaximumWidth(400)
            self.toggle_btn.setText("‹")
            self.toggle_btn.setStyleSheet(self._toggle_btn_style_expanded)
            self.select_btn.show()
            self.clear_btn.show()
 
    #TREE  
    def populate_tree(self):
        for lvl1, sub1 in self.selected_tables.items():
            p1 = QTreeWidgetItem(self.tree, [lvl1])
            p1.setFlags(p1.flags() | Qt.ItemIsUserCheckable)
            p1.setCheckState(0, Qt.Unchecked)

            for lvl2, sub2 in sub1.items():
                p2 = QTreeWidgetItem(p1, [lvl2])
                p2.setFlags(p2.flags() | Qt.ItemIsUserCheckable)
                p2.setCheckState(0, Qt.Unchecked)

                for lvl3, data in sub2.items():
                    p3 = QTreeWidgetItem(p2, [lvl3])
                    p3.setFlags(p3.flags() | Qt.ItemIsUserCheckable)
                    p3.setCheckState(0, Qt.Unchecked)
                    p3.setData(0, Qt.UserRole, data)

        self.tree.expandAll()

        self.tree.resizeColumnToContents(0)
        self.tree.updateGeometry()

    #SELECT / CLEAR ALL  
    def select_all_items(self):
        self.tree.blockSignals(True)
        root = self.tree.invisibleRootItem()
        for i in range(root.childCount()):
            item = root.child(i)
            item.setCheckState(0, Qt.Checked)
            self.update_children(item, Qt.Checked)
        self.tree.blockSignals(False)

    def clear_all_items(self):
        self.tree.blockSignals(True)
        root = self.tree.invisibleRootItem()
        for i in range(root.childCount()):
            item = root.child(i)
            item.setCheckState(0, Qt.Unchecked)
            self.update_children(item, Qt.Unchecked)
        self.tree.blockSignals(False)

    #TABLE PREVIEW
    def load_selected_table(self, item, column):
        data = item.data(0, Qt.UserRole)

        if not isinstance(data, dict):
            return

        columns = data.get("columns", [])
        rows = data.get("rows", [])

        if not columns or not rows:
            return

        row_data = rows[0]

        self.preview.clear()
        self.preview.setColumnCount(2)
        self.preview.setRowCount(len(columns))
        self.preview.setHorizontalHeaderLabels(["Parameter", "Value"])

        for i, header in enumerate(columns):
            key_item = QTableWidgetItem(str(header))
            key_item.setForeground(Qt.black)
            key_item.setBackground(QColor("#ffffff"))

            val = row_data[i] if i < len(row_data) else ""
            val_item = QTableWidgetItem(str(val))
            val_item.setForeground(Qt.black)
            val_item.setBackground(QColor("#ffffff"))

            self.preview.setItem(i, 0, key_item)
            self.preview.setItem(i, 1, val_item)

        self.preview.verticalHeader().setVisible(False)
        header = self.preview.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Fixed)

        table_width = self.preview.viewport().width()
        self.preview.setColumnWidth(0, int(table_width * 0.60))
        self.preview.setColumnWidth(1, int(table_width * 0.40))

    #CHECKBOX LOGIC
    def handle_item_changed(self, item, column):
        state = item.checkState(0)

        self.tree.blockSignals(True)

        for i in range(item.childCount()):
            child = item.child(i)
            child.setCheckState(0, state)
            self.update_children(child, state)

        self.update_parents(item)

        self.tree.blockSignals(False)

    def update_children(self, item, state):
        for i in range(item.childCount()):
            child = item.child(i)
            child.setCheckState(0, state)
            self.update_children(child, state)

    def update_parents(self, item):
        parent = item.parent()

        while parent:
            checked = 0
            unchecked = 0

            for i in range(parent.childCount()):
                st = parent.child(i).checkState(0)
                if st == Qt.Checked:
                    checked += 1
                elif st == Qt.Unchecked:
                    unchecked += 1

            if checked == parent.childCount():
                parent.setCheckState(0, Qt.Checked)
            elif unchecked == parent.childCount():
                parent.setCheckState(0, Qt.Unchecked)
            else:
                parent.setCheckState(0, Qt.PartiallyChecked)

            parent = parent.parent()

    def get_checked_tables(self):
        checked_items = []
        root = self.tree.invisibleRootItem()

        for i in range(root.childCount()):
            lvl1 = root.child(i)
            for j in range(lvl1.childCount()):
                lvl2 = lvl1.child(j)
                for k in range(lvl2.childCount()):
                    lvl3 = lvl2.child(k)
                    if lvl3.checkState(0) == Qt.Checked:
                        data = lvl3.data(0, Qt.UserRole)
                        checked_items.append((lvl3.text(0), data))

        return checked_items
  
    #EXPORT
    def export_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Excel File",
            "results.xlsx",
            "Excel Files (*.xlsx)"
        )

        if not path:
            return

        selected_tables = self.get_checked_tables()
        if not selected_tables:
            return

        GREEN_HEX   = "FF90AF13"
        WHITE_HEX   = "FFFFFFFF"
        ALT_HEX     = "FFF4F9E3"
        BORDER_HEX  = "FF90AF13"

        header_fill   = PatternFill("solid", fgColor=GREEN_HEX)
        alt_fill      = PatternFill("solid", fgColor=ALT_HEX)
        white_fill    = PatternFill("solid", fgColor=WHITE_HEX)
        header_font   = Font(bold=True, color="FFFFFFFF", name="Calibri", size=11)
        param_font    = Font(bold=True, color="FF2d2d2d", name="Calibri", size=10)
        value_font    = Font(bold=False, color="FF2d2d2d", name="Calibri", size=10)
        center_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align    = Alignment(horizontal="left",   vertical="center", wrap_text=True)

        thin_side     = Side(style="thin", color=BORDER_HEX)
        cell_border   = Border(
            left=thin_side, right=thin_side,
            top=thin_side,  bottom=thin_side
        )

        wb = openpyxl.Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        for table_name, data in selected_tables:
            columns = data.get("columns", [])
            rows    = data.get("rows",    [])

            if not columns or not rows:
                continue

            sheet_name = table_name[:31]
            ws = wb.create_sheet(title=sheet_name)

            ws.column_dimensions["A"].width = 38
            ws.column_dimensions["B"].width = 28

            for col_idx, label in enumerate(["Parameter", "Value"], start=1):
                cell = ws.cell(row=1, column=col_idx, value=label)
                cell.fill      = header_fill
                cell.font      = header_font
                cell.alignment = center_align
                cell.border    = cell_border

            ws.row_dimensions[1].height = 22

            row_data = rows[0]

            for i, param_name in enumerate(columns):
                excel_row   = i + 2
                fill_style  = alt_fill if i % 2 == 0 else white_fill

                val = row_data[i] if i < len(row_data) else ""

                #Parameter cell
                p_cell = ws.cell(row=excel_row, column=1, value=str(param_name))
                p_cell.fill      = fill_style
                p_cell.font      = param_font
                p_cell.alignment = left_align
                p_cell.border    = cell_border

                #Value cell
                v_cell = ws.cell(row=excel_row, column=2, value=str(val))
                v_cell.fill      = fill_style
                v_cell.font      = value_font
                v_cell.alignment = center_align
                v_cell.border    = cell_border

                ws.row_dimensions[excel_row].height = 18

        wb.save(path)
        self.accept()